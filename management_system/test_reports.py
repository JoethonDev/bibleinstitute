from datetime import date, timedelta
from django.test import Client, TestCase
from django.urls import reverse
from .models import AcademicYear, Course, CourseOffering, Enrollment, Grade, Quiz, Role, User, AttendanceRecord


class ReportViewTests(TestCase):
    def setUp(self):
        self.admin_role, _ = Role.objects.get_or_create(role="admin")
        self.student_role, _ = Role.objects.get_or_create(role="student")
        self.staff_role, _ = Role.objects.get_or_create(role="staff")

        self.admin = User.objects.create_user(
            username="admin", password="pass", role=self.admin_role
        )
        self.staff = User.objects.create_user(
            username="staff", password="pass", role=self.staff_role
        )
        self.student = User.objects.create_user(
            username="student", password="pass", role=self.student_role,
            study_mode="online",
        )

        self.year = AcademicYear.objects.create(
            name="2025/2026", level=1,
            starts_on=date(2025, 9, 1), ends_on=date(2026, 6, 30),
            is_current=True,
            meeting_weekdays=[0, 1, 2, 3, 4],
        )
        self.course = Course.objects.create(name="Test Course", level=1)
        self.offering = CourseOffering.objects.create(
            course=self.course, academic_year=self.year,
        )
        Enrollment.objects.create(
            student=self.student, academic_year=self.year,
            enrolled_by=self.admin,
        )

    def _login(self, user=None):
        c = Client()
        c.force_login(user or self.admin)
        return c

    def test_admin_can_access(self):
        c = self._login(self.admin)
        resp = c.get(reverse("report-dashboard"))
        self.assertEqual(resp.status_code, 200)

    def test_staff_can_access(self):
        c = self._login(self.staff)
        resp = c.get(reverse("report-dashboard"))
        self.assertEqual(resp.status_code, 200)

    def test_student_cannot_access(self):
        c = self._login(self.student)
        resp = c.get(reverse("report-dashboard"))
        self.assertNotEqual(resp.status_code, 200)

    def test_filter_by_academic_year(self):
        c = self._login(self.admin)
        resp = c.get(reverse("report-dashboard"), {"academic_year": self.year.id})
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, self.student.username)

    def test_filter_by_student(self):
        c = self._login(self.admin)
        resp = c.get(reverse("report-dashboard"), {
            "academic_year": self.year.id,
            "student": self.student.id,
        })
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, self.student.username)

    def test_filter_by_study_mode(self):
        c = self._login(self.admin)
        resp = c.get(reverse("report-dashboard"), {
            "academic_year": self.year.id,
            "study_mode": "online",
        })
        self.assertEqual(resp.status_code, 200)

    def test_csv_export(self):
        c = self._login(self.admin)
        resp = c.get(reverse("export-report-csv"), {"academic_year": self.year.id})
        self.assertEqual(resp.status_code, 200)
        self.assertIn("text/csv", resp["Content-Type"])
        self.assertIn("report_", resp["Content-Disposition"])
        self.assertIn(self.student.username, resp.content.decode())

    def test_csv_missing_year_returns_400(self):
        c = self._login(self.admin)
        resp = c.get(reverse("export-report-csv"))
        self.assertEqual(resp.status_code, 400)

    def test_csv_injection_safety(self):
        self.student.first_name = "=CMD"
        self.student.save()
        c = self._login(self.admin)
        resp = c.get(reverse("export-report-csv"), {"academic_year": self.year.id})
        body = resp.content.decode()
        self.assertIn("'=CMD", body)
        self.assertNotIn(",=CMD,", body)

    def test_xlsx_export(self):
        c = self._login(self.admin)
        resp = c.get(reverse("export-report-xlsx"), {"academic_year": self.year.id})
        self.assertEqual(resp.status_code, 200)
        self.assertIn("spreadsheetml", resp["Content-Type"])
        self.assertIn(".xlsx", resp["Content-Disposition"])

    def test_xlsx_sheet_structure(self):
        import openpyxl
        from io import BytesIO
        c = self._login(self.admin)
        resp = c.get(reverse("export-report-xlsx"), {"academic_year": self.year.id})
        wb = openpyxl.load_workbook(BytesIO(resp.content))
        self.assertIn("Grades Summary", wb.sheetnames)
        self.assertIn("Attendance Summary", wb.sheetnames)
        self.assertIn("Attendance Daily", wb.sheetnames)

    def test_xlsx_missing_year_returns_400(self):
        c = self._login(self.admin)
        resp = c.get(reverse("export-report-xlsx"))
        self.assertEqual(resp.status_code, 400)

    def test_zero_grade_denominator(self):
        other = User.objects.create_user(
            username="noattendance", password="pass", role=self.student_role,
        )
        Enrollment.objects.create(student=other, academic_year=self.year, enrolled_by=self.admin)
        c = self._login(self.admin)
        resp = c.get(reverse("export-report-csv"), {"academic_year": self.year.id})
        self.assertEqual(resp.status_code, 200)
        self.assertIn(other.username, resp.content.decode())

    def test_grade_and_attendance_in_summary(self):
        quiz = Quiz.objects.create(
            name="Test Quiz", course_offering=self.offering, course=self.course,
            total_grade=50,
        )
        Grade.objects.create(user=self.student, quiz=quiz, total_grade=40)
        for i in range(5):
            d = date(2025, 9, 1) + timedelta(days=i)
            AttendanceRecord.objects.create(
                student=self.student, academic_year=self.year,
                attendance_date=d, action="entrance",
            )
            AttendanceRecord.objects.create(
                student=self.student, academic_year=self.year,
                attendance_date=d, action="exit",
            )
        c = self._login(self.admin)
        resp = c.get(reverse("report-dashboard"), {"academic_year": self.year.id})
        self.assertContains(resp, "40")
        self.assertContains(resp, "50")
        self.assertContains(resp, "5")
