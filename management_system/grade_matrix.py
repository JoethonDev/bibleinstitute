from __future__ import annotations

from itertools import islice
from tempfile import SpooledTemporaryFile
from typing import Any

import openpyxl
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from django.core.paginator import Page
from django.db.models import Case, CharField, IntegerField, Q, Value, When
from django.db.models.functions import Concat
from django.db.models.query import QuerySet
from django.utils.translation import get_language, gettext as _

from .academic_access import READABLE_ENROLLMENT_STATUSES
from .models import AcademicYear, AcademicYearLevel, CourseOffering, Enrollment, Grade, Quiz, User
from .utils.search import normalized_contains_q, normalize_search_text


GRADE_MATRIX_PAGE_SIZE = 50
QUIZ_TYPE_ORDER = Case(
    When(quiz_type__code="weekly", then=0),
    When(quiz_type__code="final", then=1),
    default=2,
    output_field=IntegerField(),
)
COURSE_KEYS = ("weekly", "final")


def grade_matrix_selection(
    *,
    academic_year_id: int | None = None,
    level_id: int | None = None,
) -> tuple[list[AcademicYear], AcademicYear, list[AcademicYearLevel], AcademicYearLevel]:
    """Resolve one report year and level without crossing academic scopes."""
    years = list(AcademicYear.objects.order_by("-is_active", "-ordering"))
    if not years:
        raise AcademicYear.DoesNotExist

    year = next((item for item in years if item.pk == academic_year_id), None)
    if year is None:
        year = years[0]

    levels = list(
        AcademicYearLevel.objects.filter(academic_year_id=year.pk)
        .select_related("level")
        .order_by("level__ordering")
    )
    if not levels:
        raise AcademicYearLevel.DoesNotExist

    scope = next((item for item in levels if item.level_id == level_id), None)
    if scope is None:
        scope = levels[0]
    return years, year, levels, scope


def grade_matrix_student_queryset(
    scope: AcademicYearLevel,
    *,
    name: str = "",
    course_offering_ids: tuple[int, ...] = (),
) -> QuerySet[User]:
    """Return the enrolled student roster for one academic-year level."""
    if course_offering_ids:
        full_year = Q(
            enrollments__academic_year_level_id=scope.pk,
            enrollments__status__in=READABLE_ENROLLMENT_STATUSES,
            enrollments__enrollment_type=Enrollment.Type.NORMAL,
        )
        targeted = Q(
            enrollments__course_offering_id__in=course_offering_ids,
            enrollments__status__in=READABLE_ENROLLMENT_STATUSES,
        )
        enrollment_scope = full_year | targeted
    else:
        enrollment_scope = Q(
            enrollments__academic_year_level_id=scope.pk,
            enrollments__status__in=READABLE_ENROLLMENT_STATUSES,
        )

    queryset = User.objects.filter(role__role="student").filter(enrollment_scope)
    if name:
        name = normalize_search_text(name)
        queryset = queryset.annotate(
            matrix_full_name=Concat(
                "first_name", Value(" "), "last_name", output_field=CharField()
            ),
            matrix_reverse_name=Concat(
                "last_name", Value(" "), "first_name", output_field=CharField()
            ),
        )
        queryset = queryset.filter(
            normalized_contains_q(
                ("first_name", "last_name", "username", "matrix_full_name", "matrix_reverse_name"),
                name,
            )
        )
    return queryset.distinct().order_by("last_name", "first_name", "username", "pk")


def grade_matrix_courses(
    scope: AcademicYearLevel,
    *,
    course_id: int | None = None,
) -> list[dict[str, Any]]:
    """Build the ordered course-offering/quiz schema for a matrix."""
    offering_queryset = CourseOffering.objects.filter(
        academic_year_level_id=scope.pk,
    ).select_related("course").order_by("course__name", "pk")
    if course_id is not None:
        offering_queryset = offering_queryset.filter(course_id=course_id)

    offerings = list(offering_queryset)
    by_offering_id = {
        offering.pk: {
            "id": offering.course_id,
            "offering_id": offering.pk,
            "name": offering.course.name,
            "instructor": offering.instructor or offering.course.instructor or "",
            "quizzes": [],
        }
        for offering in offerings
    }
    if not by_offering_id:
        return []

    quizzes = Quiz.objects.filter(
        course_offering_id__in=tuple(by_offering_id),
    ).select_related("quiz_type").annotate(
        matrix_type_order=QUIZ_TYPE_ORDER,
    ).order_by(
        "course_offering__course__name",
        "course_offering_id",
        "matrix_type_order",
        "name",
        "pk",
    )
    for quiz in quizzes:
        quiz_type = quiz.quiz_type.code if quiz.quiz_type_id else ""
        type_label = (
            quiz.quiz_type.name_ar
            if quiz.quiz_type_id and get_language() == "ar"
            else quiz.quiz_type.name_en
            if quiz.quiz_type_id
            else str(_("Unassigned"))
        )
        by_offering_id[quiz.course_offering_id]["quizzes"].append({
            "id": quiz.pk,
            "name": quiz.name,
            "type": quiz_type,
            "type_label": type_label,
            "total": quiz.total_grade,
        })

    courses = list(by_offering_id.values())
    for course in courses:
        course["weekly_quizzes"] = [
            quiz for quiz in course["quizzes"] if quiz["type"] == "weekly"
        ]
        course["final_quizzes"] = [
            quiz for quiz in course["quizzes"] if quiz["type"] == "final"
        ]
        course["detail_span"] = len(course["quizzes"]) + len(COURSE_KEYS)
    return courses


def grade_matrix_grade_map(
    user_ids: list[int],
    quiz_ids: list[int],
) -> dict[tuple[int, int], int]:
    """Fetch one bounded grade batch for the current page or export chunk."""
    if not user_ids or not quiz_ids:
        return {}
    return {
        (row["user_id"], row["quiz_id"]): row["total_grade"]
        for row in Grade.objects.filter(
            user_id__in=user_ids,
            quiz_id__in=quiz_ids,
        ).values("user_id", "quiz_id", "total_grade")
    }


def _course_row(course: dict[str, Any], user_id: int, grades: dict[tuple[int, int], int]) -> dict[str, Any]:
    quiz_cells = []
    for quiz in course["quizzes"]:
        quiz_cells.append({
            "grade": grades.get((user_id, quiz["id"])),
            "total": quiz["total"],
            "submitted": (user_id, quiz["id"]) in grades,
        })

    totals = {}
    for key in COURSE_KEYS:
        quizzes = course[f"{key}_quizzes"]
        submitted = [
            grades[(user_id, quiz["id"])]
            for quiz in quizzes
            if (user_id, quiz["id"]) in grades
        ]
        totals[key] = {
            "earned": sum(submitted) if submitted else None,
            "total": sum(quiz["total"] for quiz in quizzes),
            "has_grade": bool(submitted),
        }
    return {
        "quiz_cells": quiz_cells,
        "weekly": totals["weekly"],
        "final": totals["final"],
    }


def grade_matrix_rows(
    students: list[User],
    courses: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    """Shape one paginated student batch without issuing per-row queries."""
    user_ids = [student.pk for student in students]
    quiz_ids = [quiz["id"] for course in courses for quiz in course["quizzes"]]
    grades = grade_matrix_grade_map(user_ids, quiz_ids)
    rows = []
    for student in students:
        rows.append({
            "id": student.pk,
            "name": student.get_full_name().strip() or student.username,
            "username": student.username,
            "courses": [
                _course_row(course, student.pk, grades)
                for course in courses
            ],
        })
    return rows


def grade_matrix_page(
    page: Page,
    courses: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    """Build rows for a database-paginated page."""
    return grade_matrix_rows(list(page.object_list), courses)


def _safe_xlsx_cell(value: Any) -> Any:
    if value is None:
        return ""
    text = str(value)
    if text[:1] in {"=", "+", "-", "@"}:
        return f"'{text}"
    return value


def _write_only_cell(sheet, value: Any, *, fill=None, font=None, alignment=None) -> WriteOnlyCell:
    cell = WriteOnlyCell(sheet, value=_safe_xlsx_cell(value))
    if fill:
        cell.fill = fill
    if font:
        cell.font = font
    if alignment:
        cell.alignment = alignment
    return cell


def grade_matrix_workbook(
    student_queryset: QuerySet[User],
    courses: list[dict[str, Any]],
) -> SpooledTemporaryFile:
    """Create a bounded-memory XLSX workbook using the same matrix schema."""
    workbook = openpyxl.Workbook(write_only=True)
    sheet = workbook.create_sheet(str(_("Grades matrix")))
    sheet.freeze_panes = "B3"
    sheet.sheet_view.showGridLines = False

    header_fill = PatternFill("solid", fgColor="E8DFCA")
    subheader_fill = PatternFill("solid", fgColor="F5EFE2")
    weekly_fill = PatternFill("solid", fgColor="E8F3F1")
    final_fill = PatternFill("solid", fgColor="FBF1DF")
    header_font = Font(bold=True, color="0D4F56")
    final_font = Font(bold=True, color="845D1D")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    top_header = [_write_only_cell(sheet, _("Student name"), fill=header_fill, font=header_font, alignment=center)]
    second_header = [_write_only_cell(sheet, "", fill=subheader_fill, font=header_font, alignment=center)]
    for course in courses:
        top_header.append(_write_only_cell(sheet, course["name"], fill=header_fill, font=header_font, alignment=center))
        top_header.extend(
            _write_only_cell(sheet, "", fill=header_fill, font=header_font, alignment=center)
            for _ in range(course["detail_span"] - 1)
        )
        for quiz in course["quizzes"]:
            fill = final_fill if quiz["type"] == "final" else weekly_fill
            font = final_font if quiz["type"] == "final" else header_font
            second_header.append(_write_only_cell(sheet, quiz["name"], fill=fill, font=font, alignment=center))
        second_header.append(_write_only_cell(sheet, _("Weekly total"), fill=weekly_fill, font=header_font, alignment=center))
        second_header.append(_write_only_cell(sheet, _("Final total"), fill=final_fill, font=final_font, alignment=center))

    sheet.append(top_header)
    sheet.append(second_header)
    sheet.column_dimensions["A"].width = 34
    for column_number in range(2, len(top_header) + 1):
        sheet.column_dimensions[get_column_letter(column_number)].width = 14

    quiz_ids = [quiz["id"] for course in courses for quiz in course["quizzes"]]
    student_iterator = student_queryset.iterator(chunk_size=500)
    while True:
        batch = list(islice(student_iterator, 500))
        if not batch:
            break
        grades = grade_matrix_grade_map([student.pk for student in batch], quiz_ids)
        for student in batch:
            row = [
                _safe_xlsx_cell(
                    f"{student.get_full_name().strip() or student.username} ({student.username} · #{student.pk})"
                )
            ]
            for course in courses:
                for quiz in course["quizzes"]:
                    grade = grades.get((student.pk, quiz["id"]))
                    row.append("—" if grade is None else f"{grade}/{quiz['total']}")
                for key in COURSE_KEYS:
                    quizzes = course[f"{key}_quizzes"]
                    submitted = [
                        grades[(student.pk, quiz["id"])]
                        for quiz in quizzes
                        if (student.pk, quiz["id"]) in grades
                    ]
                    earned = sum(submitted) if submitted else "—"
                    total = sum(quiz["total"] for quiz in quizzes)
                    row.append(f"{earned}/{total}" if total else "—")
            sheet.append(row)

    result = SpooledTemporaryFile(max_size=8 * 1024 * 1024, mode="w+b")
    workbook.save(result)
    result.seek(0)
    return result
