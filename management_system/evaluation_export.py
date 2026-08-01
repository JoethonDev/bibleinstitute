"""Bounded XLSX export for saved course-level evaluation results."""

from __future__ import annotations

import io
import re
from collections import defaultdict
from typing import TYPE_CHECKING

from django.utils.translation import gettext as _

from .models import CourseOffering, EvaluationResult, PromotionFormula, PromotionRule, QuizType

if TYPE_CHECKING:
    from .models import PromotionFormula


_FORMULA_CELL = re.compile(r"^[\s]*[=+\-@]")
_SHEET_MAX = 31


def _safe_cell(value):
    if value is None:
        return ""
    text = str(value)
    return f"'{text}" if _FORMULA_CELL.match(text) else value


def _sheet_title(raw: str, seen: set[str]) -> str:
    cleaned = re.sub(r"[\[\]*?:/\\]", "", raw).strip() or "Sheet"
    cleaned = cleaned[:_SHEET_MAX].rstrip()
    candidate = cleaned
    suffix = 2
    while candidate in seen:
        suffix_text = f" ({suffix})"
        candidate = f"{cleaned[:_SHEET_MAX - len(suffix_text)].rstrip()}{suffix_text}"
        suffix += 1
    seen.add(candidate)
    return candidate


def _results(formula, offering_id=None):
    formula_ids = [formula.pk]
    if formula.course_offering_id is None:
        formula_ids = list(PromotionFormula.objects.filter(
            academic_year_level_id=formula.academic_year_level_id,
        ).values_list("pk", flat=True))
    query = EvaluationResult.objects.filter(
        formula_id__in=formula_ids,
        course_offering__isnull=False,
    )
    if offering_id is not None:
        query = query.filter(course_offering_id=offering_id)
    return query.select_related(
        "enrollment__student",
        "course_offering__course",
        "course_offering__academic_year_level__academic_year",
        "course_offering__academic_year_level__level",
        "overridden_by",
        "formula",
    ).order_by(
        "course_offering__course__name",
        "enrollment__student__last_name",
        "enrollment__student__first_name",
        "enrollment__student__username",
    )


def _metric_label(metric, quiz_types):
    if metric.get("metric") == PromotionRule.Metric.ATTENDANCE:
        return str(_("Attendance"))
    return quiz_types.get(metric.get("quiz_type_id"), str(_("Quiz")))


def _metric_schema(formula, offering_id, quiz_types):
    schema = {}
    for result in _results(formula, offering_id).iterator(chunk_size=500):
        for metric in (result.metric_snapshot or {}).get("metrics", []):
            key = (metric.get("metric"), metric.get("quiz_type_id"))
            schema.setdefault(key, metric)
    return sorted(
        schema.values(),
        key=lambda metric: (
            metric.get("metric") == PromotionRule.Metric.ATTENDANCE,
            metric.get("quiz_type_id") or 0,
        ),
    )


def _metric_headers(schema, quiz_types):
    headers = []
    for metric in schema:
        label = _metric_label(metric, quiz_types)
        headers.extend([
            f"{label} {_('Earned')}",
            f"{label} {_('Denominator')}",
            f"{label} %",
            f"{label} {_('Min %')}",
            f"{label} {_('Passed')}",
        ])
    return headers


def _status_display(result, field):
    method = getattr(result, f"get_{field}_display", None)
    return method() if method else getattr(result, field)


def build_evaluation_workbook(formula: PromotionFormula) -> io.BytesIO:
    """Build a write-only workbook without materializing the result population."""
    import openpyxl

    workbook = openpyxl.Workbook(write_only=True)
    quiz_types = {
        item.pk: item.name_en or item.code
        for item in QuizType.objects.only("pk", "name_en", "code").iterator(chunk_size=100)
    }
    offering_ids = list(
        _results(formula).values_list("course_offering_id", flat=True).distinct().order_by("course_offering_id")
    )
    seen_titles: set[str] = set()

    if not offering_ids:
        sheet = workbook.create_sheet(_sheet_title(str(_("No Results")), seen_titles))
        sheet.append([_("No course-level evaluation results found for this formula.")])
        result = io.BytesIO()
        workbook.save(result)
        result.seek(0)
        return result

    summary = workbook.create_sheet(_sheet_title(str(_("Summary")), seen_titles))
    summary.append([
        _("Course Offering"), _("Course"), _("Student ID"), _("Student Name"),
        _("Score"), _("Status"), _("Published Lectures"),
        _("Valid Attendance Records"), _("Attendance %"),
        _("Repeat Threshold"), _("Evaluation Start"), _("Evaluation End"),
    ])
    for result_row in _results(formula).iterator(chunk_size=500):
        snapshot = result_row.metric_snapshot or {}
        attendance = next(
            (metric for metric in snapshot.get("metrics", []) if metric.get("metric") == PromotionRule.Metric.ATTENDANCE),
            {},
        )
        student = result_row.enrollment.student
        summary.append([
            _safe_cell(result_row.course_offering_id),
            _safe_cell(snapshot.get("course_name", result_row.course_offering.course.name)),
            _safe_cell(student.pk),
            _safe_cell(student.get_full_name() or student.username),
            _safe_cell(str(result_row.computed_score)),
            _safe_cell(_status_display(result_row, "computed_status")),
            _safe_cell(snapshot.get("published_lesson_count", "")),
            _safe_cell(attendance.get("earned", "")),
            _safe_cell(attendance.get("percent", "")),
            _safe_cell(result_row.formula.failed_courses_repeat_threshold),
            _safe_cell(result_row.formula.evaluation_starts_on),
            _safe_cell(result_row.formula.evaluation_ends_on),
        ])

    for offering_id in offering_ids:
        offering = CourseOffering.objects.select_related(
            "course", "academic_year_level__academic_year", "academic_year_level__level"
        ).get(pk=offering_id)
        course = offering.course
        scope = offering.academic_year_level
        sheet = workbook.create_sheet(_sheet_title(f"{course.name} - {scope.level.display_name}", seen_titles))
        schema = _metric_schema(formula, offering_id, quiz_types)
        sheet.append([
            _("Student ID"), _("Student Name"), _("Username"), _("Academic Year"),
            _("Level"), _("Course"), _("Evaluation Start"), _("Evaluation End"),
            _("Published Lectures"), *_metric_headers(schema, quiz_types),
            _("Weighted Score"), _("Computed Status"), _("Final Status"),
            _("Override Note"), _("Evaluation Errors"),
        ])
        for result_row in _results(formula, offering_id).iterator(chunk_size=500):
            snapshot = result_row.metric_snapshot or {}
            metrics = {
                (metric.get("metric"), metric.get("quiz_type_id")): metric
                for metric in snapshot.get("metrics", [])
            }
            student = result_row.enrollment.student
            row = [
                _safe_cell(student.pk),
                _safe_cell(student.get_full_name() or student.username),
                _safe_cell(student.username),
                _safe_cell(scope.academic_year.name),
                _safe_cell(scope.level.display_name),
                _safe_cell(course.name),
                _safe_cell(result_row.formula.evaluation_starts_on),
                _safe_cell(result_row.formula.evaluation_ends_on),
                _safe_cell(snapshot.get("published_lesson_count", "")),
            ]
            for schema_metric in schema:
                metric = metrics.get((schema_metric.get("metric"), schema_metric.get("quiz_type_id")), {})
                row.extend([
                    _safe_cell(metric.get("earned", "")),
                    _safe_cell(metric.get("denominator", "")),
                    _safe_cell(metric.get("percent", "")),
                    _safe_cell(metric.get("minimum_percent", "")),
                    _safe_cell(metric.get("passed", "")),
                ])
            row.extend([
                _safe_cell(str(result_row.computed_score)),
                _safe_cell(_status_display(result_row, "computed_status")),
                _safe_cell(_status_display(result_row, "final_status")),
                _safe_cell(result_row.override_note),
                _safe_cell("; ".join(snapshot.get("grading_errors", []))),
            ])
            sheet.append(row)

    result = io.BytesIO()
    workbook.save(result)
    result.seek(0)
    return result
